"""Saving cover letters as PDF, Word and plain text."""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path
from xml.sax.saxutils import escape

from .profile import Profile


def slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")[:40] or "job"


def _header_lines(profile: Profile) -> list[str]:
    contact = " | ".join(
        x for x in (profile.get("personal.email"), profile.get("personal.phone"), profile.get("personal.linkedin")) if x
    )
    location = ", ".join(x for x in (profile.get("personal.city"), profile.get("personal.state")) if x)
    return [x for x in (profile.full_name, location, contact) if x]


def save_cover_letter(text: str, profile: Profile, out_dir: Path, company: str = "", title: str = "") -> dict[str, Path]:
    """Write the letter as .pdf, .docx and .txt. Returns {format: path}."""
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = "_".join(slug(x) for x in (profile.full_name, "Cover_Letter", company, title) if x)
    date = dt.date.today().strftime("%B %d, %Y").replace(" 0", " ")
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    header = _header_lines(profile)

    paths = {
        "txt": out_dir / f"{stem}.txt",
        "docx": out_dir / f"{stem}.docx",
        "pdf": out_dir / f"{stem}.pdf",
    }
    paths["txt"].write_text(text + "\n", encoding="utf-8")
    _write_docx(paths["docx"], header, date, paragraphs)
    _write_pdf(paths["pdf"], header, date, paragraphs)
    return paths


def _write_docx(path: Path, header: list[str], date: str, paragraphs: list[str]) -> None:
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for i, line in enumerate(header):
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.bold = i == 0
        if i == 0:
            run.font.size = Pt(14)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    doc.add_paragraph(date)
    for para in paragraphs:
        doc.add_paragraph(para)
    doc.save(path)


def _write_pdf(path: Path, header: list[str], date: str, paragraphs: list[str]) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=11, leading=15, spaceAfter=10)
    name = ParagraphStyle("name", parent=body, fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=2)
    small = ParagraphStyle("small", parent=body, fontSize=10, leading=13, spaceAfter=0)

    story = []
    for i, line in enumerate(header):
        story.append(Paragraph(escape(line), name if i == 0 else small))
    story += [Spacer(1, 0.3 * inch), Paragraph(date, body)]
    for para in paragraphs:
        story.append(Paragraph(escape(para).replace("\n", "<br/>"), body))

    SimpleDocTemplate(
        str(path), pagesize=LETTER, leftMargin=inch, rightMargin=inch, topMargin=0.9 * inch, bottomMargin=inch
    ).build(story)
